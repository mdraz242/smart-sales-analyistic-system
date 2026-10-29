import io
from datetime import datetime
from typing import Dict, Any, List
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

from backend.app.services.analytics_engine import AnalyticsEngine
from backend.app.services.insights_engine import InsightsEngine
from backend.app.utils.currency import format_currency


class ExportService:
    def __init__(self, analytics_engine: AnalyticsEngine):
        self.analytics = analytics_engine
        self.insights_engine = InsightsEngine(analytics_engine)

    def generate_excel_report(
        self,
        start_utc: datetime,
        end_utc: datetime,
        prev_start_utc: datetime,
        prev_end_utc: datetime,
        currency: str = "USD"
    ) -> bytes:
        """Generates a polished multi-tab Excel spreadsheet workbook."""
        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # remove default sheet

        # 1. Fetch data
        kpis = self.analytics.get_executive_kpis(start_utc, end_utc, prev_start_utc, prev_end_utc, currency=currency)
        products = self.analytics.get_product_performance(start_utc, end_utc, prev_start_utc, prev_end_utc, limit=50)
        categories = self.analytics.get_category_analytics(start_utc, end_utc, prev_start_utc, prev_end_utc)
        insights = self.insights_engine.generate_insights(start_utc, end_utc, prev_start_utc, prev_end_utc, currency=currency)

        # Styling tokens
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Arial", size=14, bold=True, color="0F172A")
        sub_font = Font(name="Arial", size=9, italic=True, color="64748B")
        align_left = Alignment(horizontal="left", vertical="center")
        align_right = Alignment(horizontal="right", vertical="center")
        thin_border = Border(
            left=Side(style='thin', color="E2E8F0"),
            right=Side(style='thin', color="E2E8F0"),
            top=Side(style='thin', color="E2E8F0"),
            bottom=Side(style='thin', color="E2E8F0")
        )

        # TAB 1: EXECUTIVE KPIS
        ws1 = wb.create_sheet(title="Executive KPIs")
        ws1["A1"] = "Executive BI Summary & Performance KPIs"
        ws1["A1"].font = title_font
        ws1["A2"] = f"Period: {start_utc.strftime('%Y-%m-%d')} to {end_utc.strftime('%Y-%m-%d')} | Baseline: {prev_start_utc.strftime('%Y-%m-%d')} to {prev_end_utc.strftime('%Y-%m-%d')} | Currency: {currency}"
        ws1["A2"].font = sub_font

        headers1 = ["Metric", "Current Period", "Previous Period", "Delta (%)", "Trend", "Calculation Definition"]
        for col_idx, h in enumerate(headers1, 1):
            cell = ws1.cell(row=4, column=col_idx, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = align_left

        row_idx = 5
        for key, item in kpis.items():
            if isinstance(item, dict) and "formatted_current" in item:
                ws1.cell(row=row_idx, column=1, value=key.replace("_", " ").title()).alignment = align_left
                ws1.cell(row=row_idx, column=2, value=item["formatted_current"]).alignment = align_right
                ws1.cell(row=row_idx, column=3, value=item["formatted_previous"]).alignment = align_right
                ws1.cell(row=row_idx, column=4, value=f"{item['delta_pct']:+.1f}%").alignment = align_right
                ws1.cell(row=row_idx, column=5, value=item["trend"].upper()).alignment = align_left
                ws1.cell(row=row_idx, column=6, value=item["definition"]).alignment = align_left

                for c in range(1, 7):
                    ws1.cell(row=row_idx, column=c).border = thin_border
                row_idx += 1

        # TAB 2: TOP PRODUCTS
        ws2 = wb.create_sheet(title="Top Products")
        ws2["A1"] = "Product Performance Rankings & Profitability"
        ws2["A1"].font = title_font

        headers2 = ["Product Name", "SKU", "Category", "Units Sold", "Orders", "Gross Sales", "Discounts", "Refunds", "Net Sales", "% Net Sales", "Avg Price", "COGS", "Gross Profit", "Margin %", "Growth %", "Class"]
        for col_idx, h in enumerate(headers2, 1):
            cell = ws2.cell(row=3, column=col_idx, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = align_left

        for r_idx, p in enumerate(products, 4):
            ws2.cell(row=r_idx, column=1, value=p["name"]).alignment = align_left
            ws2.cell(row=r_idx, column=2, value=p["sku"]).alignment = align_left
            ws2.cell(row=r_idx, column=3, value=p["category"]).alignment = align_left
            ws2.cell(row=r_idx, column=4, value=p["units_sold"]).alignment = align_right
            ws2.cell(row=r_idx, column=5, value=p["orders_count"]).alignment = align_right
            ws2.cell(row=r_idx, column=6, value=p["gross_sales"]).alignment = align_right
            ws2.cell(row=r_idx, column=7, value=p["discounts"]).alignment = align_right
            ws2.cell(row=r_idx, column=8, value=p["refunds"]).alignment = align_right
            ws2.cell(row=r_idx, column=9, value=p["net_sales"]).alignment = align_right
            ws2.cell(row=r_idx, column=10, value=f"{p['net_sales_percentage']:.1f}%").alignment = align_right
            ws2.cell(row=r_idx, column=11, value=p["average_selling_price"]).alignment = align_right
            ws2.cell(row=r_idx, column=12, value=p["cogs"] if p["cogs"] is not None else "N/A").alignment = align_right
            ws2.cell(row=r_idx, column=13, value=p["gross_profit"] if p["gross_profit"] is not None else "N/A").alignment = align_right
            ws2.cell(row=r_idx, column=14, value=f"{p['gross_margin_pct']:.1f}%" if p["gross_margin_pct"] is not None else "N/A").alignment = align_right
            ws2.cell(row=r_idx, column=15, value=f"{p['growth_pct']:+.1f}%").alignment = align_right
            ws2.cell(row=r_idx, column=16, value=p["classification"]).alignment = align_left

            for c in range(1, 17):
                ws2.cell(row=r_idx, column=c).border = thin_border

        # TAB 3: CATEGORY PERFORMANCE
        ws3 = wb.create_sheet(title="Category Intelligence")
        ws3["A1"] = "Category Revenue & Contribution"
        ws3["A1"].font = title_font

        headers3 = ["Category Name", "Net Sales", "Share (%)", "Units Sold", "Orders", "AOV", "Growth (%)"]
        for col_idx, h in enumerate(headers3, 1):
            cell = ws3.cell(row=3, column=col_idx, value=h)
            cell.fill = header_fill
            cell.font = header_font

        for r_idx, c in enumerate(categories, 4):
            ws3.cell(row=r_idx, column=1, value=c["name"])
            ws3.cell(row=r_idx, column=2, value=c["net_sales"])
            ws3.cell(row=r_idx, column=3, value=f"{c['sales_contribution_pct']:.1f}%")
            ws3.cell(row=r_idx, column=4, value=c["units_sold"])
            ws3.cell(row=r_idx, column=5, value=c["orders_count"])
            ws3.cell(row=r_idx, column=6, value=c["aov"])
            ws3.cell(row=r_idx, column=7, value=f"{c['growth_pct']:+.1f}%")
            for col in range(1, 8):
                ws3.cell(row=r_idx, column=col).border = thin_border

        # TAB 4: AUTOMATED INSIGHTS
        ws4 = wb.create_sheet(title="Actionable Insights")
        ws4["A1"] = "Automated Business Insights & Prescriptions"
        ws4["A1"].font = title_font

        headers4 = ["Priority", "Insight Title", "Category", "Explanation", "Suggested Business Action"]
        for col_idx, h in enumerate(headers4, 1):
            cell = ws4.cell(row=3, column=col_idx, value=h)
            cell.fill = header_fill
            cell.font = header_font

        for r_idx, ins in enumerate(insights, 4):
            ws4.cell(row=r_idx, column=1, value=ins["severity"])
            ws4.cell(row=r_idx, column=2, value=ins["title"])
            ws4.cell(row=r_idx, column=3, value=ins["category"])
            ws4.cell(row=r_idx, column=4, value=ins["explanation"])
            ws4.cell(row=r_idx, column=5, value=ins["suggested_action"])
            for col in range(1, 6):
                ws4.cell(row=r_idx, column=col).border = thin_border

        # Auto-adjust column widths
        for sheet in wb.worksheets:
            for col in sheet.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                sheet.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 55)

        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)
        return stream.getvalue()

    def generate_pdf_report(
        self,
        start_utc: datetime,
        end_utc: datetime,
        prev_start_utc: datetime,
        prev_end_utc: datetime,
        currency: str = "USD"
    ) -> bytes:
        """Generates a professional executive PDF summary report using ReportLab."""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            rightMargin=36,
            leftMargin=36,
            topMargin=36,
            bottomMargin=36
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'ReportTitle',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=20,
            leading=24,
            textColor=colors.HexColor('#0F172A'),
            spaceAfter=4
        )
        subtitle_style = ParagraphStyle(
            'ReportSub',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9,
            leading=12,
            textColor=colors.HexColor('#64748B'),
            spaceAfter=15
        )
        h2_style = ParagraphStyle(
            'SectionH2',
            parent=styles['Heading2'],
            fontName='Helvetica-Bold',
            fontSize=13,
            leading=16,
            textColor=colors.HexColor('#1E293B'),
            spaceBefore=12,
            spaceAfter=6
        )
        body_style = ParagraphStyle(
            'BodyTextCustom',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9,
            leading=12,
            textColor=colors.HexColor('#334155')
        )

        elements = []

        # Header Title
        elements.append(Paragraph("WooCommerce Executive Sales Intelligence Report", title_style))
        elements.append(Paragraph(
            f"Period: {start_utc.strftime('%b %d, %Y')} – {end_utc.strftime('%b %d, %Y')} | "
            f"Prior Baseline: {prev_start_utc.strftime('%b %d, %Y')} – {prev_end_utc.strftime('%b %d, %Y')} | "
            f"Currency: {currency}",
            subtitle_style
        ))

        # KPI Table
        kpis = self.analytics.get_executive_kpis(start_utc, end_utc, prev_start_utc, prev_end_utc, currency=currency)

        elements.append(Paragraph("1. Executive Financial & Operational KPIs", h2_style))

        kpi_table_data = [
            ["Metric", "Current Period", "Previous Period", "Delta (%)", "Trend"],
            ["Net Sales", kpis["net_sales"]["formatted_current"], kpis["net_sales"]["formatted_previous"], f"{kpis['net_sales']['delta_pct']:+.1f}%", kpis["net_sales"]["trend"].upper()],
            ["Gross Sales", kpis["gross_sales"]["formatted_current"], kpis["gross_sales"]["formatted_previous"], f"{kpis['gross_sales']['delta_pct']:+.1f}%", kpis["gross_sales"]["trend"].upper()],
            ["Total Orders", kpis["total_orders"]["formatted_current"], kpis["total_orders"]["formatted_previous"], f"{kpis['total_orders']['delta_pct']:+.1f}%", kpis["total_orders"]["trend"].upper()],
            ["Units Sold", kpis["units_sold"]["formatted_current"], kpis["units_sold"]["formatted_previous"], f"{kpis['units_sold']['delta_pct']:+.1f}%", kpis["units_sold"]["trend"].upper()],
            ["Average Order Value", kpis["average_order_value"]["formatted_current"], kpis["average_order_value"]["formatted_previous"], f"{kpis['average_order_value']['delta_pct']:+.1f}%", kpis["average_order_value"]["trend"].upper()],
            ["Refund Value", kpis["total_refunds"]["formatted_current"], kpis["total_refunds"]["formatted_previous"], f"{kpis['total_refunds']['delta_pct']:+.1f}%", kpis["total_refunds"]["trend"].upper()],
            ["Refund Rate", kpis["refund_rate"]["formatted_current"], kpis["refund_rate"]["formatted_previous"], f"{kpis['refund_rate']['delta_pct']:+.1f}%", kpis["refund_rate"]["trend"].upper()],
        ]

        t = Table(kpi_table_data, colWidths=[160, 95, 95, 75, 75])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E293B')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 9),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
            ('ALIGN', (1, 0), (-1, -1), 'CENTER'),
        ]))
        elements.append(t)
        elements.append(Spacer(1, 10))

        # Top 5 Products Table
        elements.append(Paragraph("2. Leading Product Lines (By Net Sales)", h2_style))
        products = self.analytics.get_product_performance(start_utc, end_utc, prev_start_utc, prev_end_utc, rank_by="net_sales", limit=5)

        prod_table_data = [["Product Name", "SKU", "Units", "Net Sales", "Margin %", "Class"]]
        for p in products:
            margin = f"{p['gross_margin_pct']:.1f}%" if p.get('gross_margin_pct') is not None else "N/A"
            prod_table_data.append([
                p["name"][:30],
                p["sku"],
                str(p["units_sold"]),
                format_currency(p["net_sales"], currency),
                margin,
                p["classification"]
            ])

        tp = Table(prod_table_data, colWidths=[180, 85, 45, 80, 60, 80])
        tp.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#334155')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 8),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
            ('TOPPADDING', (0, 0), (-1, -1), 3),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ]))
        elements.append(tp)
        elements.append(Spacer(1, 10))

        # Insights
        elements.append(Paragraph("3. Prescriptive Business Insights", h2_style))
        insights = self.insights_engine.generate_insights(start_utc, end_utc, prev_start_utc, prev_end_utc, currency=currency)

        for ins in insights[:4]:
            p_text = f"<b>[{ins['severity']}] {ins['title']}</b><br/>{ins['explanation']}<br/><i>Action: {ins['suggested_action']}</i>"
            elements.append(Paragraph(p_text, body_style))
            elements.append(Spacer(1, 6))

        doc.build(elements)
        buffer.seek(0)
        return buffer.getvalue()
